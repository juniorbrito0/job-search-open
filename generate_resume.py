"""
The resume engine.

All resume content lives in profile/resume.json (written during setup). This
file only knows how to lay it out as a clean Word document.

    .venv/bin/python generate_resume.py

builds the master resume into Resume/:
  - "<Name> - Resume (2-page).docx"   the main version
  - "<Name> - Resume (1-page).docx"   a short version
  - "<Name> - Resume.md"              a plain-text copy (generated, do not edit)

Tailored resumes for a specific job are built by scripts/build_resume.py from a
small overlay file in tailor/, which rewords and reorders what is already in
resume.json. Nothing here ever invents experience.
"""

from __future__ import annotations

import copy
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
RESUME_FILE = ROOT / "profile" / "resume.json"

# ── Palette (the accent comes from resume.json -> accent_color) ──
INK = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x6B, 0x72, 0x80)
BODY = RGBColor(0x1F, 0x29, 0x37)
PROSE = RGBColor(0x37, 0x41, 0x51)
DEFAULT_ACCENT = "0f766e"

PAGE_WIDTH = 8.5


class _Style:
    """Per-render layout state (spacing scale, right tab stop, accent)."""

    def __init__(self, one_page: bool, accent: str):
        self.one_page = one_page
        self.sp = 0.42 if one_page else 1.0
        self.body_line = 208 if one_page else 276
        self.side_inches = 0.55 if one_page else 0.75
        self.tab_right = Inches(PAGE_WIDTH - 2 * self.side_inches)
        try:
            self.accent = RGBColor.from_string((accent or DEFAULT_ACCENT).strip("#").upper())
        except ValueError:
            self.accent = RGBColor.from_string(DEFAULT_ACCENT.upper())


# ══════════════════════════════════════════════════════
# Data
# ══════════════════════════════════════════════════════

def load_data(path: Path | None = None) -> dict:
    path = path or RESUME_FILE
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
    except FileNotFoundError:
        raise SystemExit(
            "There is no resume on file yet (profile/resume.json is missing). "
            "Finish the setup interview first, or ask Claude to import your resume."
        )
    except json.JSONDecodeError as exc:
        raise SystemExit(
            f"profile/resume.json has a typo in it (line {exc.lineno}) and could not be read. "
            "Ask Claude to fix it."
        )
    data.setdefault("experience", [])
    data.setdefault("education", [])
    data.setdefault("certifications", [])
    data.setdefault("skills_groups", [])
    data.setdefault("tools", "")
    data.setdefault("languages", [])
    if isinstance(data.get("education"), dict):
        data["education"] = [data["education"]]
    return data


def apply_overlay(data: dict, overlay: dict) -> dict:
    """Return a copy of the master data with a tailoring overlay applied.

    Overlay keys: headline, target_roles, summary, bullets {company: [..]},
    skills_line, extra_keywords [..]. Anything missing keeps the master version.
    """
    out = copy.deepcopy(data)
    for key in ("headline", "target_roles", "summary"):
        if overlay.get(key):
            out[key] = overlay[key]
    bullets = overlay.get("bullets") or {}
    if bullets:
        known = {e.get("company") for e in out["experience"]}
        unknown = [c for c in bullets if c not in known]
        if unknown:
            raise SystemExit(
                "The tailoring file names a company that is not on the master resume: "
                + ", ".join(unknown)
                + ". Company names in 'bullets' must match profile/resume.json exactly."
            )
        for entry in out["experience"]:
            if entry.get("company") in bullets:
                entry["bullets"] = list(bullets[entry["company"]])
    if overlay.get("skills_line"):
        out["skills_line"] = overlay["skills_line"]
    extra = [k for k in overlay.get("extra_keywords") or [] if k]
    if extra:
        # Posting words already proven elsewhere on the master resume
        # (added by scripts/check_resume_keywords.py --fix).
        target = "skills_line" if out.get("skills_line") else "tools"
        current = out.get(target) or ""
        add = [k for k in extra if k.lower() not in current.lower()]
        if add:
            out[target] = "  ·  ".join([x for x in [current.strip()] if x] + add)
    return out


def master_text(data: dict | None = None) -> str:
    """Every word on the master resume, for keyword checks."""
    d = data or load_data()
    chunks = [d.get("headline") or "", d.get("target_roles") or "", d.get("summary") or "",
              d.get("tools") or "", d.get("skills_line") or ""]
    for group in d.get("skills_groups") or []:
        chunks.extend(str(x) for x in group)
    for entry in d.get("experience") or []:
        chunks.append(entry.get("company") or "")
        for role in entry.get("roles") or []:
            chunks.append(str(role[0]) if role else "")
        chunks.append(entry.get("previously") or "")
        chunks.extend(entry.get("bullets") or [])
    for ed in d.get("education") or []:
        chunks.extend([ed.get("school") or "", ed.get("detail") or ""])
    for cert in d.get("certifications") or []:
        chunks.append(cert if isinstance(cert, str) else " ".join(str(v) for v in cert.values()))
    return "\n".join(chunks)


def docx_text(path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


# ══════════════════════════════════════════════════════
# DOCX rendering
# ══════════════════════════════════════════════════════

def font(run, size, bold=False, italic=False, color=INK):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def spacing(st, para, before=0, after=0, line=None):
    pPr = para._p.get_or_add_pPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), str(int(before * st.sp)))
    sp.set(qn("w:after"), str(int(after * st.sp)))
    if line:
        sp.set(qn("w:line"), str(line))
        sp.set(qn("w:lineRule"), "auto")
    pPr.append(sp)


def bottom_border(para, color="E5E7EB", size="4"):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), size)
    bot.set(qn("w:space"), "4")
    bot.set(qn("w:color"), color)
    pBdr.append(bot)
    pPr.append(pBdr)


def right_tab(st, para):
    para.paragraph_format.tab_stops.add_tab_stop(st.tab_right, WD_TAB_ALIGNMENT.RIGHT)


def add_section_label(st, doc, text):
    p = doc.add_paragraph()
    bottom_border(p)  # pBdr must precede spacing in the CT_PPr child order
    spacing(st, p, before=140, after=60)
    font(p.add_run(text.upper()), 7.5, bold=True, color=st.accent)
    p.paragraph_format.keep_with_next = True  # never strand a heading at the foot of a page
    return p


def add_bullet(st, doc, text):
    p = doc.add_paragraph()
    spacing(st, p, before=0, after=36)
    fmt = p.paragraph_format
    fmt.left_indent = Inches(0.18)
    fmt.first_line_indent = Inches(-0.18)
    font(p.add_run("•\t" + text), 9, color=BODY)
    return p


def add_company(st, doc, entry, first):
    p = doc.add_paragraph()
    spacing(st, p, before=(60 if first else 100), after=20)
    right_tab(st, p)
    font(p.add_run(entry.get("company") or ""), 10, bold=True, color=INK)
    if entry.get("location"):
        font(p.add_run(f"  ·  {entry['location']}"), 9.5, color=MUTED)
    if entry.get("dates"):
        font(p.add_run(f"\t{entry['dates']}"), 8.5, color=MUTED)

    roles = [r for r in entry.get("roles") or [] if r]
    has_prev = bool(entry.get("previously"))
    for i, role in enumerate(roles):
        title = role[0]
        date = role[1] if len(role) > 1 else None
        is_last = i == len(roles) - 1
        after = 0 if (has_prev or not is_last) else 40
        p = doc.add_paragraph()
        spacing(st, p, before=16, after=after)
        right_tab(st, p)
        font(p.add_run(title), 9, color=st.accent)
        if date:
            font(p.add_run(f"\t{date}"), 8.5, color=MUTED)

    if has_prev:
        p = doc.add_paragraph()
        spacing(st, p, before=6, after=40)
        font(p.add_run(f"Previously: {entry['previously']}"), 8, italic=True, color=MUTED)

    for b in entry.get("bullets") or []:
        add_bullet(st, doc, b)


def add_skills(st, doc, data):
    skills_line = data.get("skills_line")
    if skills_line:
        p = doc.add_paragraph()
        spacing(st, p, before=60, after=0, line=st.body_line)
        font(p.add_run(skills_line), 9, color=BODY)
    else:
        for i, group in enumerate(data.get("skills_groups") or []):
            label, items = (group + ["", ""])[:2] if isinstance(group, list) else ("", str(group))
            p = doc.add_paragraph()
            spacing(st, p, before=(60 if i == 0 else 34), after=0, line=st.body_line)
            if label:
                font(p.add_run(f"{label}:  "), 8.5, bold=True, color=st.accent)
            font(p.add_run(items), 9, color=BODY)

    if data.get("tools"):
        p = doc.add_paragraph()
        spacing(st, p, before=44 if not skills_line else 30, after=0, line=st.body_line)
        if not skills_line:
            font(p.add_run("Tools:  "), 8.5, bold=True, color=MUTED)
        font(p.add_run(data["tools"]), 8.5, color=MUTED)


def create_resume(data: dict, output_path, one_page: bool = False, headline: str | None = None):
    """Write one resume .docx from resume data (master or overlaid)."""
    st = _Style(one_page, data.get("accent_color") or DEFAULT_ACCENT)
    doc = Document()

    top_margin = Inches(0.4 if one_page else 0.65)
    side_margin = Inches(st.side_inches)
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = top_margin
        sec.left_margin = sec.right_margin = side_margin

    ns = doc.styles["Normal"]
    ns.font.name = "Calibri"
    ns.font.size = Pt(9.5)
    if one_page:
        ns.paragraph_format.line_spacing = 0.95
        ns.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    spacing(st, p, before=0, after=30)
    font(p.add_run(data.get("name") or ""), 26, bold=True, color=INK)

    headline = headline or data.get("headline") or ""
    if headline:
        p = doc.add_paragraph()
        spacing(st, p, before=0, after=22)
        font(p.add_run(headline.upper()), 8, bold=True, color=st.accent)

    if data.get("target_roles"):
        p = doc.add_paragraph()
        spacing(st, p, before=0, after=60)
        font(p.add_run(data["target_roles"]), 8, color=MUTED)

    p = doc.add_paragraph()
    spacing(st, p, before=0, after=140)
    font(p.add_run(data.get("contact") or ""), 8.5, color=MUTED)

    if data.get("summary"):
        add_section_label(st, doc, "Summary")
        p = doc.add_paragraph()
        spacing(st, p, before=0, after=0, line=st.body_line)
        font(p.add_run(data["summary"]), 9.2, color=PROSE)

    if data.get("experience"):
        add_section_label(st, doc, "Experience")
        first = True
        for entry in data["experience"]:
            if one_page and entry.get("skip_one_page"):
                continue
            add_company(st, doc, entry, first)
            first = False

    if data.get("education"):
        add_section_label(st, doc, "Education")
        for i, ed in enumerate(data["education"]):
            p = doc.add_paragraph()
            spacing(st, p, before=60 if i == 0 else 30, after=0)
            right_tab(st, p)
            font(p.add_run(ed.get("school") or ""), 9.5, bold=True, color=INK)
            if ed.get("dates"):
                font(p.add_run(f"\t{ed['dates']}"), 8.5, color=MUTED)
            if ed.get("detail"):
                p = doc.add_paragraph()
                spacing(st, p, before=0, after=0)
                font(p.add_run(ed["detail"]), 8.5, color=MUTED)

    certs = [c for c in data.get("certifications") or [] if c]
    if certs:
        add_section_label(st, doc, "Certifications")
        for i, cert in enumerate(certs):
            text = cert if isinstance(cert, str) else "  ·  ".join(
                str(cert.get(k)) for k in ("name", "issuer", "year") if cert.get(k))
            p = doc.add_paragraph()
            spacing(st, p, before=60 if i == 0 else 20, after=0)
            font(p.add_run(text), 9, color=BODY)

    if data.get("skills_groups") or data.get("skills_line") or data.get("tools"):
        add_section_label(st, doc, "Skills")
        add_skills(st, doc, data)

    langs = [l for l in data.get("languages") or [] if l]
    if langs:
        add_section_label(st, doc, "Languages")
        p = doc.add_paragraph()
        spacing(st, p, before=60, after=0)
        for i, item in enumerate(langs):
            lang, level = (list(item) + [""])[:2] if not isinstance(item, str) else (item, "")
            if i:
                font(p.add_run("  ·  "), 9, color=MUTED)
            font(p.add_run(lang), 9, bold=True, color=INK)
            if level:
                font(p.add_run(f" ({level})"), 9, color=MUTED)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Saved: {output_path}")


# ══════════════════════════════════════════════════════
# Markdown rendering
# ══════════════════════════════════════════════════════

def _md(text):
    return (text or "").replace("  ·  ", " · ")


def write_markdown(data: dict, output_path):
    L = [f"# {(data.get('name') or '').upper()}"]
    if data.get("headline"):
        L.append(f"**{_md(data['headline'])}**")
    if data.get("target_roles"):
        L.append(_md(data["target_roles"]))
    L += ["", _md(data.get("contact")), "", "---", ""]
    if data.get("summary"):
        L += ["## Summary", "", data["summary"], "", "---", ""]
    L.append("## Experience")
    for entry in data.get("experience") or []:
        L += ["", f"**{entry.get('company')}** · {entry.get('location') or ''} · *{entry.get('dates') or ''}*", ""]
        for role in entry.get("roles") or []:
            date = role[1] if len(role) > 1 else None
            L.append(f"*{role[0]}*" + (f" · {date}" if date else ""))
        if entry.get("previously"):
            L.append(f"*Previously: {_md(entry['previously'])}*")
        L.append("")
        L += [f"- {b}" for b in entry.get("bullets") or []]
        L += ["", "---"]
    if data.get("education"):
        L += ["", "## Education", ""]
        for ed in data["education"]:
            L.append(f"**{ed.get('school')}** · {ed.get('dates') or ''}")
            if ed.get("detail"):
                L.append(_md(ed["detail"]))
            L.append("")
        L.append("---")
    certs = [c for c in data.get("certifications") or [] if c]
    if certs:
        L += ["", "## Certifications", ""]
        for c in certs:
            L.append("- " + (c if isinstance(c, str) else " · ".join(
                str(c.get(k)) for k in ("name", "issuer", "year") if c.get(k))))
        L += ["", "---"]
    L += ["", "## Skills", ""]
    if data.get("skills_line"):
        L += [_md(data["skills_line"]), ""]
    else:
        for group in data.get("skills_groups") or []:
            label, items = (list(group) + ["", ""])[:2]
            L += [f"**{label}:** {_md(items)}" if label else _md(items), ""]
    if data.get("tools"):
        L += [f"*Tools: {_md(data['tools'])}*", ""]
    langs = [l for l in data.get("languages") or [] if l]
    if langs:
        L += ["---", "", "## Languages", ""]
        L.append(" · ".join(
            f"{x[0]} ({x[1]})" if not isinstance(x, str) and len(x) > 1 and x[1] else (x if isinstance(x, str) else x[0])
            for x in langs))
        L.append("")
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    Path(output_path).write_text("\n".join(L), encoding="utf-8")
    print(f"Saved: {output_path}")


# ══════════════════════════════════════════════════════

def safe_filename(text: str) -> str:
    bad = '/\\:*?"<>|'
    return "".join("-" if ch in bad else ch for ch in (text or "")).strip() or "Resume"


def main():
    data = load_data()
    name = safe_filename(data.get("name") or "Resume")
    out = ROOT / "Resume"
    create_resume(data, out / f"{name} - Resume (2-page).docx", one_page=False)
    create_resume(data, out / f"{name} - Resume (1-page).docx", one_page=True)
    write_markdown(data, out / f"{name} - Resume.md")


if __name__ == "__main__":
    sys.exit(main())
